import os
import sys
from pptx import Presentation
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from pptx.dml.color import RGBColor

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Set a solid fill first to see if it changes slide.background properties
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(255, 0, 0)

def _hex(color):
    if hasattr(color, "rgb"):
        return str(color.rgb)
    return color

def _set_gradient_bg(slide, color1="000000", color2="FFFFFF", angle=270):
    ang = angle * 60000
    c1, c2 = _hex(color1), _hex(color2)
    
    print(f"slide.background type: {type(slide.background)}")
    if hasattr(slide.background, '_element'):
        print(f"slide.background._element type: {type(slide.background._element)}")
    
    cSld = slide.background._element if hasattr(slide.background, '_element') else slide.background._cSld
    print(f"cSld type: {type(cSld)}")
    
    bg = cSld.find(qn('p:bg'))
    if bg is None:
        bg = parse_xml(f'<p:bg {nsdecls("p")}/>')
        cSld.insert(0, bg)
        
    bgPr = bg.find(qn('p:bgPr'))
    if bgPr is None:
        bgPr = parse_xml(f'<p:bgPr {nsdecls("p")}/>')
        bg.append(bgPr)
        
    fill_tags = [
        qn('a:noFill'), qn('a:solidFill'), qn('a:gradFill'), 
        qn('a:pattFill'), qn('a:grpFill'), qn('a:blipFill')
    ]
    for tag in fill_tags:
        el = bgPr.find(tag)
        if el is not None:
            bgPr.remove(el)
            
    fill_xml = (
        f'<a:gradFill {nsdecls("a")} rotWithShape="0">'
        f'  <a:gsLst>'
        f'    <a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>'
        f'    <a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>'
        f'  </a:gsLst>'
        f'  <a:lin ang="{ang}" scaled="0"/>'
        f'</a:gradFill>'
    )
    grad_elem = parse_xml(fill_xml)
    bgPr.append(grad_elem)

try:
    _set_gradient_bg(slide)
    print("Success")
except Exception as e:
    import traceback
    traceback.print_exc()
