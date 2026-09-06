import os
import sys
from pptx import Presentation
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])

shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(255, 0, 0)

def _set_shape_transparency(shape, alpha_val):
    print("shape.fill:", type(shape.fill))
    print("shape.fill._fill:", type(shape.fill._fill))
    if hasattr(shape.fill, '_xPr'):
        print("shape.fill._xPr:", type(shape.fill._xPr))
    
    # Try different ways to find a:solidFill
    try:
        solid_fill = shape.fill._fill
        srgb = solid_fill.find(qn('a:solidFill'))
        print("Success finding a:solidFill")
    except Exception as e:
        print(f"Error calling find on _fill: {e}")
        
    try:
        spPr = shape.fill._xPr # this is usually the properties element
        solidFill = spPr.find(qn('a:solidFill'))
        if solidFill is not None:
            srgbClr = solidFill.find(qn('a:srgbClr'))
            if srgbClr is not None:
                alpha_elem = parse_xml(f'<a:alpha {nsdecls("a")} val="{alpha_val}"/>')
                srgbClr.append(alpha_elem)
                print("Successfully added alpha using _xPr")
    except Exception as e:
        print(f"Error using _xPr: {e}")

_set_shape_transparency(shape, 50000)
